"""Day 2 regression tests: no external AI requests or real Chroma needed."""
from __future__ import annotations

import io
import sqlite3

import pytest
from openpyxl import load_workbook

from app import services
from app.reports import build_report, parse_report_sections


class FakeStore:
    def __init__(self, fail=False):
        self.fail = fail
        self.adds = []
        self.deletes = []

    def add_documents(self, chunks, ids):
        self.adds.append((chunks, ids))
        if self.fail:
            raise RuntimeError("Simulated embedding failure")

    def delete(self, ids):
        self.deletes.append(ids)


@pytest.fixture
def isolated_db(tmp_path, monkeypatch):
    fake = FakeStore()
    monkeypatch.setattr(services, "DB_FILE", tmp_path / "documents.sqlite3")
    monkeypatch.setattr(services, "vector_store", lambda: fake)
    return fake


def test_second_upload_skips_embeddings(isolated_db):
    first = services.index_document(b"Gloves and goggles required.", "policy.txt")
    second = services.index_document(b"Gloves and goggles required.", "renamed.txt")
    assert first["duplicate"] is False
    assert second["duplicate"] is True
    assert first["id"] == second["id"]
    assert second["filename"] == "policy.txt"  # Existing original document.
    assert len(services.list_documents()) == 1
    assert len(isolated_db.adds) == 1


def test_different_bytes_indexed_separately(isolated_db):
    a = services.index_document(b"Wear gloves.", "policy.txt")
    b = services.index_document(b"Wear goggles.", "policy.txt")
    assert a["id"] != b["id"]
    assert len(services.list_documents()) == 2
    assert len(isolated_db.adds) == 2


def test_embedding_failure_does_not_record_document(tmp_path, monkeypatch):
    fake = FakeStore(fail=True)
    monkeypatch.setattr(services, "DB_FILE", tmp_path / "documents.sqlite3")
    monkeypatch.setattr(services, "vector_store", lambda: fake)
    with pytest.raises(RuntimeError, match="Simulated embedding failure"):
        services.index_document(b"Incident report.", "report.txt")
    assert services.list_documents() == []
    assert len(fake.deletes) == 1


def test_legacy_database_migrates_without_deleting_rows(tmp_path, monkeypatch):
    db = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(db) as con:
        con.execute("""CREATE TABLE documents (id TEXT PRIMARY KEY, filename TEXT NOT NULL,
            chunks INTEGER NOT NULL, uploaded_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""")
        con.execute("INSERT INTO documents(id, filename, chunks) VALUES('legacy', 'old.txt', 1)")
    monkeypatch.setattr(services, "DB_FILE", db)
    assert services.list_documents()[0]["id"] == "legacy"
    with sqlite3.connect(db) as con:
        fields = {row[1] for row in con.execute("PRAGMA table_info(documents)")}
    assert "sha256" in fields


def test_structured_excel_export():
    answer = """Preliminary Gap Review
(1) Documented Controls
- Goggles are required [1].
(2) Potential Gaps or Unanswered Questions
- Worker said no goggles were used [2].
- Training status is not documented [2].
(3) Follow-up Actions
- Check PPE training records [1][2].
Summary: Follow-up requires verification.
"""
    parts = parse_report_sections(answer)
    assert len(parts["Potential gaps"]) == 2
    data, *_ = build_report("Review", answer, [
        {"marker": "[1]", "filename": "policy.txt", "locator": "text", "excerpt": "Goggles required"},
        {"marker": "[2]", "filename": "incident.txt", "locator": "text", "excerpt": "No goggles"},
    ], "xlsx")
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Analysis", "Findings", "Follow-up Actions", "Evidence"]
    assert book["Findings"].max_row == 3
    assert book["Findings"]["B2"].value == "[2]"
    assert book["Follow-up Actions"].max_row == 2
    assert book["Evidence"].max_row == 3
    assert book["Analysis"]["B2"].value.startswith("AI-generated draft")


def test_spreadsheet_formula_protection():
    content, *_ = build_report("Review", "=HYPERLINK(\"https://bad.test\",\"x\")", [
        {"marker": "[1]", "filename": "demo.txt", "locator": "text", "excerpt": "=1+1"}
    ], "xlsx")
    book = load_workbook(io.BytesIO(content))
    assert book["Analysis"]["B5"].data_type != "f"
    assert book["Evidence"]["D2"].data_type != "f"