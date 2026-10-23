"""Regression tests for grouped source references in Excel exports.

No OpenAI API calls, ChromaDB, or saved document database are needed.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.reports import build_report, citation_audit, extract_citation_markers


SOURCES = [
    {"marker": "[1]", "filename": "demo_policy.txt", "locator": "text",
     "excerpt": "Workers must use protective equipment."},
    {"marker": "[2]", "filename": "demo_policy.txt", "locator": "text",
     "excerpt": "Corrective actions require an owner and due date."},
    {"marker": "[3]", "filename": "demo_incident.txt", "locator": "text",
     "excerpt": "Worker reported not wearing safety goggles."},
]


def test_extracts_grouped_and_individual_markers_without_duplicates():
    text = "Gap found [1, 3]; follow-up [2][3]; repeated [1,3]."
    assert extract_citation_markers(text) == ["[1]", "[3]", "[2]"]


def test_grouped_citations_satisfy_policy_and_incident_audit():
    audit = citation_audit("PPE discrepancy [1, 3].", SOURCES, comparison=True)
    assert audit["all"] == "[1], [3]"
    assert audit["policy"] == "[1]"
    assert audit["incident"] == "[3]"
    assert audit["status"] == "Markers valid — verify excerpt support"


def test_unknown_marker_stays_invalid_and_is_never_added_as_evidence():
    audit = citation_audit("Verify record [1, 99]", SOURCES, comparison=False)
    assert audit["all"] == "[1]"
    assert "Unrecognized markers: [99]" in audit["status"]


def test_excel_export_recognizes_grouped_citations_in_both_sheets():
    answer = """(1) Documented Controls
- Goggles are required [1].
(2) Potential Gaps or Unanswered Questions
- Goggles were not worn, contrary to policy [1, 3].
(3) Follow-up Actions
- Review PPE supervision and training records [1, 3].
Summary: Professional review required.
"""
    payload, mime, ext = build_report("EHS Gap Review", answer, SOURCES, "xlsx")
    assert ext == "xlsx" and "spreadsheetml" in mime
    book = load_workbook(BytesIO(payload), read_only=True)
    assert book.sheetnames == ["Analysis", "Findings", "Follow-up Actions", "Evidence"]

    findings = book["Findings"]
    assert findings["B2"].value == "[1], [3]"  # All citations
    assert findings["C2"].value == "[1]"       # Policy evidence
    assert findings["D2"].value == "[3]"       # Incident evidence
    assert findings["F2"].value == "Markers valid — verify excerpt support"

    actions = book["Follow-up Actions"]
    assert actions["B2"].value == "[1], [3]"
    assert actions["C2"].value == "Markers valid — verify excerpt support"
    book.close()
