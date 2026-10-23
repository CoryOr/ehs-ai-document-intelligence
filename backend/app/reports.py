# reports.py
"""Export sourced EHS reviews as PDF, Word, and structured Excel workbooks."""
from __future__ import annotations

import re
from html import escape
from io import BytesIO

from docx import Document as WordDocument
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

DISCLAIMER = "AI-generated draft for demonstration only. Requires professional EHS review."


def excel_text(text: str) -> str:
    """Keep untrusted file excerpts and LLM text from becoming Excel formulas."""
    value = str(text)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def parse_report_sections(answer: str) -> dict[str, list[str]]:
    """Split the existing numbered LLM review into rows without inventing facts.

    This only understands the format requested by answer_question(review=True).
    Unstructured answers remain available in the Analysis sheet.
    """
    sections: dict[str, list[str]] = {
        "Overview": [],
        "Documented controls": [],
        "Potential gaps": [],
        "Follow-up actions": [],
        "Summary": [],
    }
    section = "Overview"
    headers = {
        "documented controls": "Documented controls",
        "potential gaps or unanswered questions": "Potential gaps",
        "follow-up actions": "Follow-up actions",
    }
    for raw in answer.splitlines():
        line = raw.strip()
        if not line:
            continue
        header = re.match(r"^\(?[123]\)?[.)\s-]*(.+?)\s*$", line, re.I)
        if header and header.group(1).strip().lower() in headers:
            section = headers[header.group(1).strip().lower()]
            continue
        if line.lower().startswith("summary:"):
            section = "Summary"
            line = line.split(":", 1)[1].strip()
        if line.startswith(("- ", "• ", "* ")):
            sections[section].append(line[2:].strip())
        elif section not in ("Overview", "Summary") and sections[section]:
            # A wrapped continuation line belongs to the preceding item.
            sections[section][-1] += " " + line
        else:
            sections[section].append(line)
    return sections


def source_role(source: dict) -> str:
    """Best-effort source classification using document names, NOT a verified fact.

    Users may upload arbitrary filenames. Unrecognized names remain ``other``
    rather than being incorrectly called a policy or an incident.
    """
    name = str(source.get("filename", "")).casefold().replace("-", "_")
    if "policy" in name or "procedure" in name:
        return "policy"
    if "incident" in name or "near_miss" in name:
        return "incident"
    return "other"


# Bracketed numeric citation groups, e.g. [1], [1, 3], or [2,3, 5].
# Ignore arbitrary bracketed prose such as [policy] rather than guessing.
_CITATION_GROUP = re.compile(r"\[\s*(\d+(?:\s*,\s*\d+)*)\s*\]")


def extract_citation_markers(text: str) -> list[str]:
    """Normalize grouped and individual references in first-seen order.

    "[1, 3] and [3][2]" becomes ["[1]", "[3]", "[2]"].
    A marker's presence does not prove its excerpt supports a claim.
    """
    markers: list[str] = []
    seen: set[str] = set()
    for group in _CITATION_GROUP.finditer(text):
        for item in group.group(1).split(","):
            marker = f"[{int(item.strip())}]"
            if marker not in seen:
                seen.add(marker)
                markers.append(marker)
    return markers


def citation_audit(text: str, sources: list[dict], *, comparison: bool) -> dict:
    """Check citation marker existence and document-type coverage, NOT claim truth.

    Never attach a source marker the model did not provide. Marker presence is
    structural evidence only; a human must still verify actual semantic support.
    """
    by_marker = {str(s.get("marker", "")): s for s in sources}
    claimed = extract_citation_markers(text)
    valid = [mark for mark in claimed if mark in by_marker]
    invalid = [mark for mark in claimed if mark not in by_marker]
    policy = [mark for mark in valid if source_role(by_marker[mark]) == "policy"]
    incident = [mark for mark in valid if source_role(by_marker[mark]) == "incident"]
    other = [mark for mark in valid if source_role(by_marker[mark]) == "other"]

    problems = []
    if invalid:
        problems.append("Unrecognized markers: " + ", ".join(invalid))
    if not valid:
        problems.append("No valid supporting citations")
    if comparison:
        available_roles = {source_role(s) for s in sources}
        for role in ("policy", "incident"):
            if role not in available_roles:
                problems.append(f"No {role} source retrieved")
            elif not (policy if role == "policy" else incident):
                problems.append(f"Missing {role} citation")
    return {
        "all": ", ".join(valid),
        "policy": ", ".join(policy),
        "incident": ", ".join(incident),
        "other": ", ".join(other),
        "status": "; ".join(problems) if problems else "Markers valid — verify excerpt support",
    }


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="153D57")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = ws.dimensions


def _wrap_rows(ws, from_row: int = 2) -> None:
    for row in ws.iter_rows(min_row=from_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row[0].row].height = 64


def build_report(title: str, answer: str, sources: list[dict], format: str) -> tuple[bytes, str, str]:
    if format == "docx":
        doc = WordDocument()
        doc.add_heading(title, 0)
        doc.add_paragraph(DISCLAIMER)
        for paragraph in answer.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph)
        doc.add_heading("Retrieved evidence", level=1)
        for s in sources:
            doc.add_heading(f"{s['marker']} {s['filename']} — {s['locator']}", level=2)
            doc.add_paragraph(s["excerpt"])
        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"

    if format == "pdf":
        buffer = BytesIO()
        styles = getSampleStyleSheet()
        pdf = SimpleDocTemplate(buffer, pagesize=letter, title=title)
        story = [Paragraph(escape(title), styles["Title"]),
                 Paragraph(escape(DISCLAIMER), styles["Italic"]), Spacer(1, 16)]
        for line in answer.split("\n"):
            if line.strip():
                story.append(Paragraph(escape(line), styles["BodyText"]))
                story.append(Spacer(1, 7))
        story.append(Spacer(1, 12))
        story.append(Paragraph("Retrieved evidence", styles["Heading2"]))
        for s in sources:
            story.append(Paragraph(escape(f"{s['marker']} {s['filename']} — {s['locator']}"), styles["Heading3"]))
            story.append(Paragraph(escape(s["excerpt"]), styles["BodyText"]))
            story.append(Spacer(1, 8))
        pdf.build(story)
        return buffer.getvalue(), "application/pdf", "pdf"

    if format == "xlsx":
        parts = parse_report_sections(answer)
        book = Workbook()
        overview = book.active
        overview.title = "Analysis"
        overview.append(["EHS Intelligence Report", excel_text(title)])
        overview.append(["Notice", DISCLAIMER])
        overview.append([])
        overview.append(["Section", "Review text"])
        for section, lines in parts.items():
            for line in lines:
                overview.append([section, excel_text(line)])
        for cell in overview[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="153D57")
        overview["A1"].font = Font(bold=True, size=13, color="153D57")
        overview["B1"].font = Font(bold=True, size=13, color="153D57")
        overview.column_dimensions["A"].width = 26
        overview.column_dimensions["B"].width = 105
        overview.freeze_panes = "B5"
        overview.auto_filter.ref = f"A4:B{max(4, overview.max_row)}"
        _wrap_rows(overview, 5)
        overview["B2"].alignment = Alignment(wrap_text=True)

        findings = book.create_sheet("Findings")
        findings.append([
            "Potential gap / unanswered question", "Citations", "Policy evidence",
            "Incident evidence", "Other evidence", "Citation check", "Review status", "Reviewer notes",
        ])
        for gap in parts["Potential gaps"]:
            audit = citation_audit(gap, sources, comparison=True)
            findings.append([
                excel_text(gap), audit["all"], audit["policy"], audit["incident"],
                audit["other"], audit["status"], "Needs human review", "",
            ])
        _header(findings)
        for col, width in {"A": 88, "B": 18, "C": 18, "D": 18, "E": 18,
                           "F": 52, "G": 25, "H": 42}.items():
            findings.column_dimensions[col].width = width
        _wrap_rows(findings)

        actions = book.create_sheet("Follow-up Actions")
        actions.append(["Suggested follow-up", "Citations", "Evidence check", "Owner", "Progress"])
        for action in parts["Follow-up actions"]:
            audit = citation_audit(action, sources, comparison=False)
            actions.append([excel_text(action), audit["all"], audit["status"], "", "Not started"])
        _header(actions)
        for col, width in {"A": 88, "B": 20, "C": 52, "D": 26, "E": 23}.items():
            actions.column_dimensions[col].width = width
        _wrap_rows(actions)

        evidence = book.create_sheet("Evidence")
        evidence.append(["Citation", "Document", "Location", "Retrieved excerpt"])
        for s in sources:
            evidence.append([excel_text(s["marker"]), excel_text(s["filename"]),
                             excel_text(s["locator"]), excel_text(s["excerpt"])])
        _header(evidence)
        for col, width in {"A": 15, "B": 32, "C": 25, "D": 95}.items():
            evidence.column_dimensions[col].width = width
        _wrap_rows(evidence)

        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    raise ValueError("Unsupported report format")